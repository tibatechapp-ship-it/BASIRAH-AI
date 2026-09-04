"""اختبار دعم ملفات DOCX."""
from pathlib import Path

import pytest

from basirah.io.file_handler import extract_text, read_docx


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    """إنشاء ملف DOCX تجريبي للاختبار."""
    from docx import Document
    
    doc = Document()
    doc.add_heading('تجربة ملف Word', 0)
    doc.add_paragraph('هذا نص تجريبي باللغة العربية لاختبار دعم ملفات DOCX.')
    doc.add_paragraph('البصيرة هو مشروع لتصنيف الكتب الإسلامية تلقائياً.')
    
    file_path = tmp_path / "test.docx"
    doc.save(str(file_path))
    return file_path


def test_read_docx_basic(sample_docx: Path):
    """اختبار قراءة ملف DOCX أساسي."""
    text = read_docx(str(sample_docx))
    
    assert text != ""
    assert len(text) > 0
    assert 'تجريبي' in text
    assert 'البصيرة' in text


def test_read_docx_nonexistent():
    """اختبار قراءة ملف غير موجود."""
    text = read_docx("/tmp/nonexistent_file.docx")
    assert text == ""


def test_extract_text_docx(sample_docx: Path):
    """اختبار استخراج النص من ملف DOCX عبر extract_text."""
    text = extract_text(str(sample_docx))
    
    assert text != ""
    assert 'تجريبي' in text
    assert 'البصيرة' in text


def test_extract_text_unsupported_extension(tmp_path: Path):
    """اختبار امتداد ملف غير مدعوم."""
    unsupported_file = tmp_path / "test.xyz"
    unsupported_file.write_text("test", encoding="utf-8")
    
    text = extract_text(str(unsupported_file))
    assert text == ""


def test_read_docx_empty_paragraphs(tmp_path: Path):
    """اختبار ملف DOCX يحتوي على فقرات فارغة فقط."""
    from docx import Document
    
    doc = Document()
    # إضافة فقرات فارغة فقط
    doc.add_paragraph("")
    doc.add_paragraph("   ")
    doc.add_paragraph("\t")
    
    file_path = tmp_path / "empty.docx"
    doc.save(str(file_path))
    
    text = read_docx(str(file_path))
    # يجب أن يعيد سلسلة فارغة أو نص فارغ
    assert text.strip() == ""


def test_read_docx_mixed_content(tmp_path: Path):
    """اختبار ملف DOCX يحتوي على محتوى مختلط (عربي وإنجليزي)."""
    from docx import Document
    
    doc = Document()
    doc.add_heading('Mixed Content Test', 0)
    doc.add_paragraph('This is English text.')
    doc.add_paragraph('هذا نص عربي.')
    doc.add_paragraph('123456789')
    
    file_path = tmp_path / "mixed.docx"
    doc.save(str(file_path))
    
    text = read_docx(str(file_path))
    
    assert 'English' in text
    assert 'عربي' in text
    assert '123456789' in text
